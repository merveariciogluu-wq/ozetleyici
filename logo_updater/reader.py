# -*- coding: utf-8 -*-
"""
reader.py — Open a source .xlsx file, detect yellow/blue colored cells,
and return a list of ColoredCell named-tuples.
"""

import sys
import os
from collections import namedtuple

import openpyxl
from openpyxl.styles import PatternFill, GradientFill

# Add parent dir to path so we can import config when run standalone
sys.path.insert(0, os.path.dirname(__file__))
from config import SOURCE_SHEET_NAME, YELLOW_RGB, PRODUCT_CODE_COL

ColoredCell = namedtuple("ColoredCell", [
    "row_idx",       # 0-based index within data rows (row 2 in file = index 0)
    "col_idx",       # 0-based column index
    "value",         # the cell's value
    "color",         # "yellow" | "blue"
    "product_code",  # value from column J of the same row
    "row_values",    # tuple of all cell values in the row
])


def _rgb_is_blue(rgb_str: str) -> bool:
    """Return True if the AARRGGBB hex string represents a blue-dominant color."""
    if not rgb_str or len(rgb_str) < 6:
        return False
    try:
        # AARRGGBB format
        if len(rgb_str) == 8:
            rr = int(rgb_str[2:4], 16)
            gg = int(rgb_str[4:6], 16)
            bb = int(rgb_str[6:8], 16)
        elif len(rgb_str) == 6:
            rr = int(rgb_str[0:2], 16)
            gg = int(rgb_str[2:4], 16)
            bb = int(rgb_str[4:6], 16)
        else:
            return False
        return bb > rr and bb > gg
    except ValueError:
        return False


def _check_fg_color(fg) -> str:
    """
    Given an fgColor object, return 'yellow', 'blue', or ''.
    Works regardless of whether fg comes from a direct PatternFill
    or a StyleProxy wrapper.
    """
    if fg is None:
        return ""
    try:
        color_type = fg.type
    except Exception:
        return ""

    if color_type == "rgb":
        rgb = fg.rgb
        if not rgb or rgb == "00000000":
            return ""
        if rgb.upper() == YELLOW_RGB:
            return "yellow"
        if _rgb_is_blue(rgb):
            return "blue"

    elif color_type == "theme":
        try:
            theme_idx = fg.theme
            tint = fg.tint if fg.tint is not None else 0.0
        except Exception:
            return ""
        if theme_idx in (4, 5):
            return "blue"
        if theme_idx is not None and theme_idx >= 8 and tint > 0:
            return "blue"

    return ""


def _is_yellow(fill) -> bool:
    """Return True if fill is a solid yellow (FFFFFF00) pattern fill."""
    if not hasattr(fill, "fgColor"):
        return False
    return _check_fg_color(fill.fgColor) == "yellow"


def _is_blue(fill) -> bool:
    """
    Return True if the fill represents a blue background.
    Handles PatternFill (direct or via StyleProxy) and GradientFill.
    """
    # Try PatternFill-style: has fgColor
    if hasattr(fill, "fgColor"):
        result = _check_fg_color(fill.fgColor)
        if result == "blue":
            return True
        return False

    # Try GradientFill-style: has stop_list or stop
    for attr in ("stop", "stop_list"):
        stops = getattr(fill, attr, None)
        if stops:
            try:
                for stop in stops:
                    color = getattr(stop, "color", None)
                    if color and _check_fg_color(color) == "blue":
                        return True
            except (TypeError, AttributeError):
                pass

    return False


def _get_cell_color(cell) -> str:
    """Return 'yellow', 'blue', or '' for a cell."""
    # MergedCell objects may not have a fill attribute
    if not hasattr(cell, "fill") or cell.fill is None:
        return ""
    fill = cell.fill
    if _is_yellow(fill):
        return "yellow"
    if _is_blue(fill):
        return "blue"
    return ""


def read_colored_cells(filepath: str) -> list:
    """
    Open filepath (data_only=False) and return a list of ColoredCell for every
    cell that has a yellow or blue background fill.

    Raises:
        ValueError: if the "Rapor" sheet is not found in the workbook.
        Exception: for any other file-level error.

    Rules:
    - Only data rows (row 2 onward) are processed.
    - Rows where column J (Ürün Kodu) is None or empty are skipped.
    - Each flagged cell yields one ColoredCell entry.
    """
    wb = openpyxl.load_workbook(filepath, data_only=False, read_only=False)

    if SOURCE_SHEET_NAME not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"'{SOURCE_SHEET_NAME}' sheet not found in '{filepath}'. "
            f"Available sheets: {wb.sheetnames}"
        )

    ws = wb[SOURCE_SHEET_NAME]
    results = []
    row_idx = 0  # 0-based counter for data rows

    for row_cells in ws.iter_rows(min_row=2):
        values = tuple(cell.value for cell in row_cells)

        # Determine product code from column J (index 9, 0-based)
        product_code = None
        if PRODUCT_CODE_COL < len(values):
            product_code = values[PRODUCT_CODE_COL]

        # Skip rows with no product code
        if product_code is None or str(product_code).strip() == "":
            row_idx += 1
            continue

        # Check each cell for color
        for col_idx, cell in enumerate(row_cells):
            color = _get_cell_color(cell)
            if color:
                cell_value = values[col_idx] if col_idx < len(values) else None
                results.append(ColoredCell(
                    row_idx=row_idx,
                    col_idx=col_idx,
                    value=cell_value,
                    color=color,
                    product_code=product_code,
                    row_values=values,
                ))

        row_idx += 1

    wb.close()
    return results
